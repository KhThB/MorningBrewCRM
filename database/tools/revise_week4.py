from pathlib import Path
import ast, math, json
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

QA=Path('database/qa/tuan4/revised');QA.mkdir(parents=True,exist_ok=True)
OUT=Path('database/docs/tuan4/1.docx')
F='C:/Windows/Fonts/arial.ttf';FB='C:/Windows/Fonts/arialbd.ttf'
# Chỉ dùng lại hàm vẽ, không chạy bộ sinh báo cáo cũ.
tree=ast.parse(Path('database/tools/build_bfd_dfd.py').read_text(encoding='utf-8'))
defs=[n for n in tree.body if isinstance(n,ast.FunctionDef) and n.name in ['font','wrap','text','node','arrow','band']]
exec(compile(ast.Module(body=defs,type_ignores=[]),'<drawing helpers>','exec'))
N={'E1':'Khách hàng','E2':'Nhân viên','E3':'Cấp quản lý','D1':'Hồ sơ khách hàng','D2':'Sản phẩm và giá','D3':'Đơn hàng','D4':'Sổ điểm thành viên','D5':'Ưu đãi khách hàng','D6':'Phản hồi','D7':'Khảo sát'}
P={1:'Bán hàng tại chi nhánh',2:'Quản lý điểm thành viên',3:'Phân nhóm và triển khai ưu đãi',4:'Tiếp nhận và xử lý phản hồi',5:'Tổ chức khảo sát khách hàng',6:'Báo cáo và thống kê'}
C={
1:[
 ('Nhận diện khách hàng thành viên',[('E2','Thông tin nhận diện khách'),('D1','Hồ sơ thành viên')],[('E2','Kết quả nhận diện khách')]),
 ('Lập đơn và tính tiền theo chi nhánh',[('E2','Sản phẩm, số lượng và chi nhánh'),('D2','Sản phẩm và giá chi nhánh')],[('E2','Chi tiết đơn và tiền hàng')]),
 ('Kiểm tra ưu đãi áp dụng',[('E2','Thông tin ưu đãi xuất trình'),('D5','Ưu đãi và tình trạng sử dụng')],[('E2','Kết quả áp dụng và tiền phải trả')]),
 ('Ghi nhận thanh toán và hoàn tất đơn',[('E2','Xác nhận thu tiền, phương thức')],[('D3','Đơn hoàn tất và chi tiết'),('D5','Thông tin ưu đãi đã sử dụng'),('E2','Kết quả ghi nhận giao dịch')]),
],
2:[
 ('Xác định điểm từ đơn hoàn tất',[('D3','Đơn hoàn tất của thành viên'),('D4','Điểm đã ghi theo đơn')],[]),
 ('Ghi nhận biến động điểm',[('E3','Yêu cầu điều chỉnh điểm có lý do'),('D1','Khách hàng và số dư'),('D4','Lịch sử biến động điểm')],[('D4','Giao dịch cộng hoặc điều chỉnh'),('D1','Số dư điểm mới'),('E3','Kết quả điều chỉnh điểm')]),
 ('Tra cứu điểm và lịch sử biến động',[('E1','Yêu cầu xem điểm cá nhân'),('D1','Số dư điểm khách hàng'),('D4','Các giao dịch điểm')],[('E1','Số dư và lịch sử điểm')]),
],
3:[
 ('Phân tích hồ sơ và hành vi mua',[('E3','Tiêu chí phân nhóm'),('D1','Tuổi, giới tính và sở thích'),('D3','Lịch sử, tần suất và giá trị mua')],[]),
 ('Xác định nhóm khách phù hợp',[],[('E3','Danh sách khách theo tiêu chí')]),
 ('Xét điều kiện và cấp ưu đãi',[('E3','Chương trình và yêu cầu cấp'),('E1','Yêu cầu sử dụng điểm nhận ưu đãi'),('D1','Khách hàng và số dư điểm'),('D4','Lịch sử sử dụng điểm'),('D5','Điều kiện chương trình')],[('D5','Ưu đãi cấp cho khách'),('D4','Giao dịch trừ điểm nếu có'),('D1','Số dư sau sử dụng điểm'),('E3','Kết quả cấp ưu đãi')]),
 ('Gửi và tra cứu ưu đãi',[('E1','Yêu cầu xem ưu đãi'),('D5','Ưu đãi đã cấp và trạng thái')],[('E1','Ưu đãi và điều kiện sử dụng')]),
],
4:[
 ('Tiếp nhận phản hồi',[('E1','Phản hồi, đánh giá, quán hoặc món'),('D1','Danh tính khách gửi')],[]),
 ('Phân loại và xác định nơi xử lý',[],[('D6','Phản hồi phân loại theo phạm vi'),('E3','Phản hồi thuộc phạm vi xử lý')]),
 ('Xử lý và trả lời khách',[('E3','Kết quả xử lý và nội dung trả lời'),('D6','Phản hồi chờ giải quyết')],[('D6','Trạng thái xử lý và câu trả lời'),('E1','Trả lời chính thức')]),
 ('Theo dõi kết quả giải quyết',[('E1','Yêu cầu xem phản hồi cá nhân'),('E3','Điều kiện theo dõi phản hồi'),('D6','Phản hồi và kết quả giải quyết')],[('E1','Tình trạng phản hồi'),('E3','Danh sách và tiến độ giải quyết')]),
],
5:[
 ('Thiết lập nội dung khảo sát',[('E3','Chủ đề, câu hỏi và thời hạn')],[('D7','Khảo sát và bộ câu hỏi')]),
 ('Chọn khách hàng nhận khảo sát',[('E3','Tiêu chí chọn người nhận'),('D1','Hồ sơ và sở thích'),('D3','Lịch sử mua trong phạm vi')],[('D7','Danh sách người nhận khảo sát')]),
 ('Gửi khảo sát đến khách hàng',[('E3','Yêu cầu phát hành khảo sát'),('D7','Khảo sát và danh sách nhận')],[('E1','Nội dung và lời mời khảo sát'),('D7','Thông tin gửi khảo sát')]),
 ('Kiểm tra và ghi nhận câu trả lời',[('E1','Câu trả lời khảo sát'),('D7','Câu hỏi, thời hạn và người nhận')],[('D7','Câu trả lời và trạng thái hoàn thành'),('E1','Kết quả gửi câu trả lời')]),
 ('Tổng hợp kết quả khảo sát',[('E3','Yêu cầu tổng hợp khảo sát'),('D7','Câu hỏi và câu trả lời')],[('E3','Kết quả tổng hợp từng khảo sát')]),
],
6:[
 ('Thống kê khách hàng',[('E3','Bộ lọc và phạm vi khách hàng'),('D1','Hồ sơ và sở thích'),('D3','Tần suất, lịch sử mua')],[('E3','Thống kê và cơ cấu khách hàng')]),
 ('Thống kê doanh thu và đơn hàng',[('E3','Kỳ và phạm vi báo cáo bán hàng'),('D3','Đơn hoàn tất và chi tiết')],[('E3','Doanh thu và số lượng đơn')]),
 ('Thống kê điểm và ưu đãi',[('E3','Điều kiện báo cáo thành viên'),('D4','Giao dịch điểm'),('D5','Ưu đãi đã cấp và sử dụng')],[('E3','Báo cáo điểm và ưu đãi')]),
 ('Báo cáo phản hồi và khảo sát',[('E3','Kỳ và phạm vi báo cáo CRM'),('D6','Phản hồi và kết quả xử lý'),('D7','Khảo sát và câu trả lời')],[('E3','Báo cáo chất lượng chăm sóc')]),
]}
LINKS={1:[(1,2,'Khách thành viên hoặc khách vãng lai'),(2,3,'Đơn và tiền hàng trước ưu đãi'),(3,4,'Đơn đã xác định tiền phải trả')],2:[(1,2,'Điểm được tính theo đơn')],3:[(1,2,'Hồ sơ và chỉ tiêu mua đã phân tích'),(2,3,'Nhóm khách đủ điều kiện')],4:[(1,2,'Phản hồi hợp lệ')],5:[],6:[]}
SHORT={
1:({'E2':'Thông tin khách, đơn và thu tiền','D1':'Hồ sơ thành viên','D2':'Sản phẩm và giá','D5':'Điều kiện và trạng thái ưu đãi'},{'E2':'Kết quả lập đơn và giao dịch','D3':'Đơn và chi tiết hoàn tất','D5':'Ưu đãi đã sử dụng'}),
2:({'D3':'Đơn hoàn tất của thành viên','D4':'Lịch sử điểm','E3':'Yêu cầu điều chỉnh điểm','D1':'Hồ sơ và số dư','E1':'Yêu cầu tra cứu điểm'},{'D4':'Giao dịch biến động điểm','D1':'Số dư mới','E3':'Kết quả điều chỉnh','E1':'Số dư và lịch sử điểm'}),
3:({'E3':'Tiêu chí và yêu cầu cấp ưu đãi','D1':'Hồ sơ, sở thích, số dư','D3':'Lịch sử mua','E1':'Yêu cầu nhận hoặc xem ưu đãi','D4':'Lịch sử sử dụng điểm','D5':'Chương trình, ưu đãi đã cấp'},{'E3':'Nhóm khách và kết quả cấp','D5':'Ưu đãi cấp cho khách','D4':'Giao dịch sử dụng điểm','D1':'Số dư sau sử dụng điểm','E1':'Ưu đãi và điều kiện'}),
4:({'E1':'Phản hồi, yêu cầu theo dõi','D1':'Danh tính khách','E3':'Trả lời và điều kiện theo dõi','D6':'Phản hồi hiện có'},{'D6':'Phản hồi và kết quả xử lý','E3':'Phản hồi và tiến độ giải quyết','E1':'Trả lời và tình trạng phản hồi'}),
5:({'E3':'Nội dung, tiêu chí, yêu cầu khảo sát','D7':'Khảo sát và dữ liệu trả lời','D1':'Hồ sơ và sở thích','D3':'Lịch sử mua','E1':'Câu trả lời khảo sát'},{'D7':'Khảo sát, người nhận, câu trả lời','E1':'Khảo sát và kết quả gửi','E3':'Tổng hợp từng khảo sát'}),
6:({'E3':'Điều kiện và phạm vi báo cáo','D1':'Hồ sơ khách','D3':'Đơn và lịch sử mua','D4':'Giao dịch điểm','D5':'Ưu đãi cấp và sử dụng','D6':'Phản hồi đã xử lý','D7':'Khảo sát và câu trả lời'},{'E3':'Báo cáo bán hàng và CRM'}),
}
for k in P:
 for direction,idx in [(0,1),(1,2)]:
  assert set(SHORT[k][direction])=={key for s in C[k] for key,_ in s[idx]}
 for j,s in enumerate(C[k],1):
  assert s[1] or any(b==j for a,b,_ in LINKS[k])
  assert s[2] or any(a==j for a,b,_ in LINKS[k])
im=Image.new('RGB',(3600,1690),'white');d=ImageDraw.Draw(im)
node(d,(1100,25,1400,120),'E','MORNING BREW CRM',45)
d.line([(1800,145),(1800,210)],fill='black',width=3);d.line([(290,210),(3290,210)],fill='black',width=3)
for k in P:
 x=30+(k-1)*600;d.line([(x+260,210),(x+260,260)],fill='black',width=3)
 node(d,(x,260,520,150),'E',f'{k}.0\n{P[k]}',33)
 d.line([(x+25,410),(x+25,480+(len(C[k])-1)*240+75)],fill='black',width=3)
 for j,(name,_,_) in enumerate(C[k],1):
  yy=465+(j-1)*240;d.line([(x+25,yy+80),(x+70,yy+80)],fill='black',width=3)
  node(d,(x+70,yy,475,160),'E',f'{k}.{j}\n{name}',31)
im.save(QA/'bfd.png')
im=Image.new('RGB',(2000,1040),'white');d=ImageDraw.Draw(im)
band(d,30,70,1940,'0','Hệ thống\nMorning Brew CRM',[
 ('E1','Yêu cầu điểm, ưu đãi, phản hồi và câu trả lời khảo sát'),('E2','Thông tin khách, đơn hàng, ưu đãi và xác nhận thu tiền'),('E3','Tiêu chí CRM, điều chỉnh điểm, xử lý phản hồi, khảo sát, báo cáo')],[
 ('E1','Điểm, ưu đãi, trả lời phản hồi, khảo sát và kết quả tiếp nhận'),('E2','Kết quả nhận diện, tiền phải trả và giao dịch'),('E3','Nhóm khách, kết quả xử lý, tổng hợp khảo sát và báo cáo')],step=290,size=34)
im.save(QA/'dfd0.png')
im=Image.new('RGB',(3700,2690),'white');d=ImageDraw.Draw(im)
for k in P:
 band(d,35+((k-1)%2)*1850,35+((k-1)//2)*880,1770,f'{k}.0',P[k],list(SHORT[k][0].items()),list(SHORT[k][1].items()),step=117,size=28)
im.save(QA/'dfd1.png')
for k in P:
 heights=[max(len(s[1]),len(s[2]),2)*112 for s in C[k]]
 im=Image.new('RGB',(2400,sum(heights)+80*len(heights)+30),'white');d=ImageDraw.Draw(im);y=25
 for j,(name,ins,outs) in enumerate(C[k],1):
  h=band(d,25,y,2350,f'{k}.{j}',name,ins,outs,step=112,size=31)
  for a,b,label in LINKS[k]:
   if a==j:
    assert b==a+1
    arrow(d,(1200,y+h),(1200,y+h+80));text(d,(1240,y+h+3,710,75),label,25)
  y+=h+80
 im.save(QA/f'dfd2_{k}.png')

doc=Document();s=doc.sections[0];s.page_width,s.page_height=Cm(42),Cm(29.7)
s.top_margin=s.bottom_margin=Cm(1.3);s.left_margin=s.right_margin=Cm(1.5)
for n in ['Normal','Title','Heading 1']:
 st=doc.styles[n];st.font.name='Arial';st.font.size=Pt(12 if n=='Normal' else 20 if n=='Title' else 16);st.font.color.rgb=RGBColor(0,0,0)
 for el in list(st.element.iter(qn('w:pBdr'))):el.getparent().remove(el)
doc.styles['Normal'].paragraph_format.space_after=Pt(5)
def p(t):doc.add_paragraph(t)
def h(t):doc.add_paragraph(t,'Heading 1')
def page(t):doc.add_page_break();h(t)
def pic(name,maxh=22):
 im=Image.open(QA/name);w,hh=im.size;sc=min(39/w,maxh/hh)
 pp=doc.add_paragraph();pp.alignment=1;pp.add_run().add_picture(str(QA/name),width=Cm(w*sc),height=Cm(hh*sc))
def table(headers,rows):
 tb=doc.add_table(rows=1,cols=len(headers));tb.style='Table Grid'
 for c,t in zip(tb.rows[0].cells,headers):c.text=t
 for rr in rows:
  for c,t in zip(tb.add_row().cells,rr):c.text=t
 for i,row in enumerate(tb.rows):
  row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
  if i==0:row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
  for c in row.cells:
   cp=c._tc.get_or_add_tcPr();bd=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    x=OxmlElement('w:'+edge);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'D9D9D9');bd.append(x)
   cp.append(bd)
   if i==0:
    x=OxmlElement('w:shd');x.set(qn('w:fill'),'EEEEEE');cp.append(x)
    for r in c.paragraphs[0].runs:r.bold=True
doc.add_paragraph('SƠ ĐỒ BFD VÀ DFD MORNING BREW CRM','Title')
p('Nhóm 1 • Báo cáo tuần 4 • Ngày 01 tháng 10 năm 2026')
p('Hệ thống quản lý chuỗi quán cà phê kết hợp CRM. Báo cáo mô hình hóa bán hàng tại chi nhánh, điểm thành viên, phân nhóm và ưu đãi, phản hồi, khảo sát, báo cáo thống kê. Nhân viên lập đơn và ghi nhận thanh toán; khách hàng tham gia các nghiệp vụ thành viên và chăm sóc khách hàng.')
h('Phạm vi và ký hiệu')
p('BFD chỉ phân rã các nghiệp vụ đặc trưng, không liệt kê thêm, xóa, sửa. Xác thực, phân quyền, tổ chức chi nhánh và cấu hình sản phẩm là chức năng hỗ trợ; mọi nghiệp vụ vẫn tuân thủ quyền và phạm vi được giao.')
p('DFD mức 0 là sơ đồ ngữ cảnh; mức 1 gồm 6 tiến trình chính; mức 2 phân rã thành 24 tiến trình con. Mã tiến trình thống nhất giữa BFD và DFD. Sơ đồ mô tả dữ liệu, không dùng mũi tên để biểu diễn điều kiện hoặc thứ tự thao tác.')
p('Hình chữ nhật: tác nhân ngoài. Hình chữ nhật bo góc: tiến trình. Hình hở cạnh phải: kho dữ liệu. Mũi tên: luồng dữ liệu một chiều. Ký hiệu có cùng mã được vẽ lặp để giảm giao cắt; đó vẫn là cùng một tác nhân hoặc kho dùng chung.')
table(['Tác nhân','Vai trò nghiệp vụ'],[(f'{k} {N[k]}',v) for k,v in [('E1','Xem điểm và ưu đãi, sử dụng điểm theo chương trình, gửi phản hồi và trả lời khảo sát.'),('E2','Nhân viên thu ngân hoặc nhân viên được phân công lập đơn, xác nhận giao dịch tại chi nhánh.'),('E3','Chủ doanh nghiệp, quản lý khu vực, quản lý chi nhánh. Chủ xem toàn chuỗi; cấp dưới chỉ thao tác trong phạm vi được giao.')]])
h('Kho dữ liệu logic')
table(['Kho','Nội dung'],[
 ('D1 Hồ sơ khách hàng','Users, Customers, CustomerPreferences: thông tin nhận diện, sở thích và số dư điểm.'),
 ('D2 Sản phẩm và giá','Products, Categories, BranchProducts: sản phẩm và giá bán tại chi nhánh.'),
 ('D3 Đơn hàng','Orders, OrderDetails: đơn, phương thức thanh toán, giá tại lúc bán, chi tiết và khách thành viên nếu có.'),
 ('D4 Sổ điểm thành viên','LoyaltyTransactions: lịch sử cộng, trừ và điều chỉnh điểm.'),
 ('D5 Ưu đãi khách hàng','Promotions, CustomerPromotions: chương trình, điều kiện, người nhận và tình trạng sử dụng.'),
 ('D6 Phản hồi','Feedbacks: nội dung, loại phản hồi, chi nhánh/sản phẩm, trạng thái và trả lời chính thức.'),
 ('D7 Khảo sát','Surveys, SurveyQuestions, SurveyRecipients, SurveyResponses: nội dung, người nhận và câu trả lời.')])
p('Kho logic không tương ứng một-một với bảng. Quy tắc quy đổi điểm và điều kiện chương trình phải được thống nhất trước triển khai; sơ đồ không tự đặt mức điểm, hạn mức hay công thức giảm cụ thể.')
page('1 Sơ đồ BFD')
pic('bfd.png',20)
p('BFD gồm 6 nhóm nghiệp vụ, 24 chức năng con. Việc áp dụng và ghi nhận sử dụng ưu đãi nằm trong bán hàng; việc sử dụng điểm để nhận ưu đãi nằm trong nhóm 3.0. Không phân rã nghiệp vụ giao hàng, thanh toán trực tuyến hoặc hoàn tiền trong phạm vi báo cáo này.')
page('2 Sơ đồ DFD mức 0')
pic('dfd0.png',22)
p('Admin, Staff/POS và Customer là các giao diện của cùng hệ thống, không phải tác nhân ngoài. Khách mua tại quầy trao đổi với nhân viên; dữ liệu đơn vào hệ thống qua E2. Mức ngữ cảnh không thể hiện kho dữ liệu.')
page('3 Sơ đồ DFD mức 1')
pic('dfd1.png',23)
p('Các tiến trình dùng chung kho theo mã D1–D7. Tên gói dữ liệu ở mức 1 bao quát các thành phần chi tiết trên DFD mức 2 tương ứng.')
notes={
1:'Nhận diện thành viên hoặc khách vãng lai, tính tiền theo giá chi nhánh, kiểm tra ưu đãi, ghi nhận thanh toán và đơn hoàn tất. Đơn lưu đơn giá tại thời điểm bán. Chỉ đơn gắn thành viên mới là căn cứ tích điểm. D3 cung cấp dữ liệu cho 2.0 và các báo cáo.',
2:'Xác định điểm từ đơn hoàn tất theo quy tắc của chương trình, kiểm tra lịch sử để tránh cộng trùng. Điều chỉnh điểm cần người có quyền và lý do. Số dư D1 phải khớp lịch sử D4, bao gồm cả điểm sử dụng ở 3.3. Khách chỉ xem dữ liệu của mình.',
3:'Phân nhóm dựa trên hồ sơ và lịch sử mua. Cấp ưu đãi theo nhóm hoặc xét yêu cầu sử dụng điểm theo chương trình. Trường hợp có tiêu điểm, phải kiểm tra đủ điểm và ghi đồng thời số dư, giao dịch điểm, ưu đãi được cấp. 1.3 kiểm tra ưu đãi khi mua; 1.4 ghi nhận đã sử dụng vào D5.',
4:'Phản hồi có thể liên quan quán, món hoặc nội dung dịch vụ. Xác định loại và phạm vi để chuyển cấp quản lý phù hợp, lưu trạng thái cùng trả lời chính thức. Luồng tiếp nhận không bắt buộc phản hồi phải gắn với một đơn hàng.',
5:'Khảo sát được xác định nội dung, câu hỏi và thời hạn; người nhận được chọn theo tiêu chí hồ sơ hoặc lịch sử mua. Chỉ ghi nhận trả lời phù hợp người nhận, câu hỏi và thời gian mở; theo dõi trạng thái hoàn thành. 5.5 tổng hợp từng khảo sát, còn 6.4 lập báo cáo tổng hợp theo kỳ và phạm vi.',
6:'Các báo cáo lọc theo thời gian, khu vực hoặc chi nhánh tùy quyền. Doanh thu dựa trên đơn hoàn tất; điểm và ưu đãi lấy từ sổ giao dịch và trạng thái cấp/sử dụng. Báo cáo phản hồi, khảo sát phục vụ đánh giá chăm sóc khách hàng.'}
for k in P:
 page(f'4.{k} DFD mức 2 của tiến trình {k}.0 {P[k]}')
 pic(f'dfd2_{k}.png',22)
 p('Diễn giải nghiệp vụ: '+notes[k])
page('5 Đối chiếu phân rã và cân bằng dữ liệu')
table(['Tiến trình mức 1','Phân rã mức 2','Biên dữ liệu'],[
 (f'{k}.0 {P[k]}',', '.join(f'{k}.{i}' for i in range(1,len(C[k])+1)), 'Vào: '+', '.join(SHORT[k][0])+'. Ra: '+', '.join(SHORT[k][1])+'.') for k in P])
h('Các luồng nội bộ khi phân rã')
table(['Luồng','Dữ liệu'],[(f'{k}.{a} → {k}.{b}',label) for k in P for a,b,label in LINKS[k]])
p('Các luồng nội bộ trên chỉ xuất hiện ở mức 2 và không làm thay đổi đầu vào/ra của tiến trình cha. Các tiến trình con khác trao đổi qua kho dùng chung, ví dụ 5.1 và 5.2 ghi D7 để 5.3 gửi khảo sát; 5.4 ghi D7 để 5.5 tổng hợp.')
h('Quan hệ giữa các nhóm nghiệp vụ')
p('1.0 ghi đơn hoàn tất vào D3 để 2.0 tính điểm, 3.0 phân tích khách và 6.0 thống kê. 3.0 ghi ưu đãi vào D5 để 1.0 kiểm tra khi bán. Biến động điểm từ 2.0 và 3.0 cùng lưu D4, cập nhật số dư D1. Phản hồi D6 và khảo sát D7 cung cấp dữ liệu cho 6.0.')
p('Toàn bộ luồng với E1, E2, E3 ở mức 1 nằm trong các gói dữ liệu tương ứng trên mức 0. Cấp quản lý được gộp thành E3 theo loại tương tác, nhưng không đồng nghĩa các vai trò có quyền như nhau. Thông tin phân công và quyền được kiểm tra trong mọi tiến trình liên quan.')
footer=s.footer.paragraphs[0];footer.alignment=2;footer.add_run('Nhóm 1 • Morning Brew CRM • Trang ')
x=OxmlElement('w:fldSimple');x.set(qn('w:instr'),'PAGE');footer._p.append(x)
assert sum(map(len,C.values()))==24
doc.save(OUT)
(QA/'model.json').write_text(json.dumps({'N':N,'P':P,'C':C,'LINKS':LINKS,'SHORT':SHORT},ensure_ascii=False,indent=2),encoding='utf-8')
print('Updated:',OUT.resolve());print('6 groups; 24 subfunctions; all boundary and process I/O checks passed')
