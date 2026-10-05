from pathlib import Path
from collections import defaultdict, Counter
import json, math, textwrap
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE=Path('database/docs/tuan4'); BASE.mkdir(parents=True,exist_ok=True)
QA=Path('database/qa/tuan4'); QA.mkdir(parents=True,exist_ok=True)
N={
 'E1':'Khách hàng', 'E2':'Nhân viên POS', 'E3':'Cấp quản lý',
 'D1':'Hồ sơ khách hàng', 'D2':'Thực đơn chi nhánh', 'D3':'Đơn và giao hàng',
 'D4':'Điểm và voucher', 'D5':'Thanh toán và hoàn tiền', 'D6':'Phản hồi và khảo sát',
}
P={1:'Tiếp nhận đặt món',2:'Phục vụ và xử lý đơn',3:'Thanh toán và hoàn tiền',4:'Tích điểm và cấp ưu đãi',5:'Chăm sóc khách hàng',6:'Phân tích kinh doanh'}
# Mỗi dòng là một luồng dữ liệu có hướng. Các kho lặp lại vẫn là cùng một kho logic.
C={
1:[
 ('Kiểm tra yêu cầu đặt món', [('E1','Chi nhánh, món, tùy chọn'),('E2','Thông tin đơn tại quầy'),('D1','Hồ sơ và địa chỉ khách'),('D2','Món, giá và khả năng bán')], [('E1','Menu và lỗi lựa chọn'),('E2','Menu và lỗi nhập đơn')]),
 ('Tính giá và kiểm tra voucher', [('D4','Voucher và điều kiện giảm')], [('E1','Chi tiết tiền và lỗi voucher'),('E2','Chi tiết tiền và lỗi voucher')]),
 ('Chốt đơn và giữ voucher', [], [('D3','Đơn, chi tiết và địa chỉ chốt'),('D4','Voucher giữ cho đơn'),('E1','Mã đơn và số tiền phải trả'),('E2','Mã đơn và số tiền phải trả')]),
],
2:[
 ('Tiếp nhận và chuẩn bị món', [('E2','Xác nhận đơn và tiến độ pha chế'),('D3','Đơn đang chờ phục vụ'),('D5','Kết quả thanh toán')], [('D3','Trạng thái xác nhận, chuẩn bị'),('E2','Danh sách đơn cần phục vụ'),('E1','Tình trạng chuẩn bị đơn')]),
 ('Bàn giao và theo dõi giao hàng', [('E2','Thông tin bàn giao, giao lại'),('D3','Đơn sẵn sàng và thông tin giao'),('D5','Xác nhận đã thu đủ tiền')], [('D3','Kết quả giao và hoàn tất đơn'),('D4','Voucher đã sử dụng'),('E1','Kết quả giao nhận'),('E2','Kết quả xử lý giao nhận')]),
 ('Kiểm tra và xử lý hủy đơn', [('E1','Yêu cầu hủy đơn của khách'),('E3','Yêu cầu hủy và lý do'),('D3','Trạng thái đơn và hạn thanh toán'),('D5','Tình trạng thu tiền')], [('D3','Đơn hủy và lý do'),('D4','Voucher được giải phóng'),('E1','Kết quả hủy đơn'),('E3','Kết quả xử lý hủy')]),
],
3:[
 ('Ghi nhận thanh toán', [('E1','Yêu cầu thanh toán hoặc thử lại'),('E2','Xác nhận tiền đã thu tại quầy'),('D3','Mã đơn, số tiền và hạn trả'),('D5','Các lần thanh toán trước')], [('D5','Lần thanh toán và kết quả'),('E1','Kết quả thanh toán'),('E2','Kết quả ghi nhận thu tiền')]),
 ('Xử lý hoàn tiền toàn bộ', [('E3','Yêu cầu hoàn tiền và lý do'),('D3','Đơn hủy hoặc đơn cần hoàn'),('D5','Giao dịch thu và yêu cầu hoàn')], [('D5','Yêu cầu và kết quả hoàn tiền'),('E1','Tình trạng hoàn tiền'),('E3','Kết quả xử lý hoàn tiền')]),
 ('Đối chiếu giao dịch', [('E3','Điều kiện tra cứu giao dịch'),('D3','Tổng tiền và tình trạng đơn'),('D5','Lịch sử thu và hoàn tiền')], [('E3','Danh sách giao dịch và sai lệch')]),
],
4:[
 ('Cộng và thu hồi điểm', [('D1','Khách hàng và số dư điểm'),('D3','Đơn hoàn tất và tiền hàng'),('D5','Kết quả thu và hoàn tiền'),('D4','Chính sách và sổ điểm')], [('D1','Số dư điểm sau đối chiếu'),('D4','Bút toán cộng hoặc thu hồi'),('E1','Thông tin biến động điểm')]),
 ('Đổi điểm lấy voucher', [('E1','Yêu cầu đổi điểm lấy voucher'),('D1','Số dư điểm hiện tại'),('D4','Điều kiện đổi và hạn mức')], [('D1','Số dư sau đổi điểm'),('D4','Bút toán đổi và voucher cấp'),('E1','Kết quả đổi điểm và voucher')]),
 ('Cấp ưu đãi theo nhóm khách', [('E3','Điều kiện ưu đãi và tiêu chí nhận'),('D1','Thông tin và sở thích khách'),('D3','Lịch sử mua trong phạm vi'),('D4','Chương trình và hạn mức cấp')], [('D4','Voucher cấp và người nhận'),('E1','Ưu đãi được cấp'),('E3','Kết quả cấp ưu đãi')]),
],
5:[
 ('Tiếp nhận và giải quyết phản hồi', [('E1','Nội dung phản hồi và đánh giá'),('E3','Trả lời chính thức của quản lý'),('D1','Danh tính khách hàng'),('D3','Đơn thuộc khách và chi nhánh'),('D6','Phản hồi đang xử lý')], [('D6','Phản hồi và trả lời chính thức'),('E1','Kết quả tiếp nhận và trả lời'),('E3','Phản hồi cần xử lý')]),
 ('Phát hành khảo sát đúng đối tượng', [('E3','Nội dung, thời hạn và tiêu chí gửi'),('D1','Hồ sơ và sở thích khách'),('D3','Lịch sử mua theo phạm vi')], [('D6','Khảo sát và danh sách được mời'),('E1','Lời mời và nội dung khảo sát'),('E3','Kết quả phát hành khảo sát')]),
 ('Tiếp nhận câu trả lời khảo sát', [('E1','Bộ câu trả lời khảo sát'),('D6','Câu hỏi, đáp án, hạn và người nhận')], [('D6','Câu trả lời đã kiểm tra'),('E1','Kết quả gửi khảo sát')]),
],
6:[
 ('Thống kê doanh thu và đơn hàng', [('E3','Kỳ và phạm vi báo cáo bán hàng'),('D3','Đơn hoàn tất và chi tiết món'),('D5','Hoàn tiền đã hoàn thành')], [('E3','Doanh thu, đơn và món bán chạy')]),
 ('Phân tích khách hàng và ưu đãi', [('E3','Tiêu chí phân tích khách hàng'),('D1','Hồ sơ và số dư điểm'),('D3','Lịch sử mua hàng'),('D4','Sổ điểm và voucher đã cấp')], [('E3','Nhóm khách và hiệu quả ưu đãi')]),
 ('Tổng hợp phản hồi và khảo sát', [('E3','Điều kiện tổng hợp chăm sóc'),('D6','Phản hồi và câu trả lời khảo sát')], [('E3','Thống kê phản hồi và khảo sát')]),
]}
# Tên ngắn của gói dữ liệu ở mức 1. Thành phần chi tiết nằm trên mức 2 tương ứng.
SHORT={
1:{'i':{'E1':'Yêu cầu đặt món','E2':'Dữ liệu đơn tại quầy','D1':'Hồ sơ, địa chỉ','D2':'Menu, giá, trạng thái','D4':'Điều kiện voucher'},'o':{'E1':'Menu, tiền và kết quả đặt','E2':'Menu, tiền và kết quả lập đơn','D3':'Hồ sơ đơn đã chốt','D4':'Voucher đang giữ'}},
2:{'i':{'E1':'Yêu cầu hủy đơn','E2':'Xác nhận, pha chế, giao nhận','E3':'Yêu cầu hủy có lý do','D3':'Đơn, tiến độ và hạn trả','D5':'Tình trạng thanh toán'},'o':{'E1':'Tiến độ và kết quả xử lý','E2':'Đơn chờ và kết quả giao','E3':'Kết quả hủy','D3':'Trạng thái đơn, giao, hủy','D4':'Voucher dùng hoặc giải phóng'}},
3:{'i':{'E1':'Yêu cầu thanh toán','E2':'Xác nhận thu tiền','E3':'Yêu cầu hoàn, đối chiếu','D3':'Đơn và số tiền','D5':'Lịch sử giao dịch'},'o':{'E1':'Kết quả thu và hoàn','E2':'Kết quả ghi nhận thu','E3':'Kết quả hoàn, đối chiếu','D5':'Giao dịch thu và hoàn'}},
4:{'i':{'E1':'Yêu cầu đổi voucher','E3':'Điều kiện cấp ưu đãi','D1':'Hồ sơ và số dư','D3':'Đơn hoàn tất, lịch sử mua','D4':'Quy tắc, sổ điểm, voucher','D5':'Kết quả thu và hoàn'},'o':{'E1':'Biến động điểm, voucher','E3':'Kết quả cấp ưu đãi','D1':'Số dư mới','D4':'Sổ điểm và voucher cấp'}},
5:{'i':{'E1':'Phản hồi, câu trả lời','E3':'Trả lời, nội dung khảo sát','D1':'Hồ sơ khách','D3':'Lịch sử đơn theo phạm vi','D6':'Phản hồi, khảo sát hiện có'},'o':{'E1':'Trả lời, khảo sát, xác nhận','E3':'Phản hồi và kết quả phát hành','D6':'Phản hồi và khảo sát xử lý'}},
6:{'i':{'E3':'Điều kiện báo cáo','D1':'Hồ sơ khách','D3':'Đơn và lịch sử mua','D4':'Sổ điểm, voucher','D5':'Hoàn tiền thành công','D6':'Phản hồi, câu trả lời'},'o':{'E3':'Báo cáo kinh doanh và CRM'}},
}
for k,subs in C.items():
 for ix,direction in [(1,'i'),(2,'o')]:
  assert set(x for sub in subs for x,_ in sub[ix])==set(SHORT[k][direction]),(k,direction)
 for j,(name,ins,outs) in enumerate(subs):
  assert ins or (k==1 and j>0)
  assert outs or (k==1 and j<2)

F='C:/Windows/Fonts/arial.ttf'; FB='C:/Windows/Fonts/arialbd.ttf'
def font(size,bold=False):return ImageFont.truetype(FB if bold else F,size)
def wrap(draw,text,ft,width):
 lines=[]
 for para in text.split('\n'):
  current=''
  for w in para.split():
   trial=(current+' '+w).strip()
   if current and draw.textlength(trial,font=ft)>width:lines.append(current);current=w
   else:current=trial
  lines.append(current)
 return lines
def text(draw,box,content,size=28,bold=False):
 x,y,w,h=box;ft=font(size,bold);lines=wrap(draw,content,ft,w-12);lh=size*1.23
 assert lh*len(lines)<=h+2,(content,box,lines)
 yy=y+(h-lh*len(lines))/2
 for line in lines:
  draw.text((x+w/2,yy),line,font=ft,fill='black',anchor='mt');yy+=lh
def node(draw,box,key,label=None,size=27):
 x,y,w,h=box
 if key.startswith('D'):
  draw.line([(x+w,y),(x,y),(x,y+h),(x+w,y+h)],fill='black',width=2)
 elif key.startswith('P'):
  draw.rounded_rectangle((x,y,x+w,y+h),radius=24,outline='black',width=3,fill='#F3F3F3')
 else:draw.rectangle((x,y,x+w,y+h),outline='black',width=2,fill='white')
 text(draw,box,label or f'{key}\n{N[key]}',size,key.startswith('P'))
def arrow(draw,a,b,label='',size=25):
 draw.line([a,b],fill='black',width=2)
 angle=math.atan2(b[1]-a[1],b[0]-a[0]);length=12
 draw.polygon([b,(b[0]-length*math.cos(angle-.4),b[1]-length*math.sin(angle-.4)),(b[0]-length*math.cos(angle+.4),b[1]-length*math.sin(angle+.4))],fill='black')
 if label:
  ft=font(size);width=abs(b[0]-a[0])-18;lines=wrap(draw,label,ft,width);lh=size*1.15
  yy=a[1]-len(lines)*lh/2
  tw=max(draw.textlength(l,font=ft) for l in lines)
  draw.rectangle(((a[0]+b[0])/2-tw/2-5,yy-2,(a[0]+b[0])/2+tw/2+5,yy+len(lines)*lh+2),fill='white')
  for l in lines:draw.text(((a[0]+b[0])/2,yy),l,font=ft,fill='black',anchor='mt');yy+=lh
def band(draw,x,y,w,code,name,ins,outs,step=104,size=28):
 # Input -> process -> output; duplicate actors/stores explicitly labelled by ID.
 n=max(len(ins),len(outs),2);h=n*step
 nw=w*.155; px=x+w*.39;pw=w*.22;rx=x+w*.845
 node(draw,(px,y,pw,h), 'P'+code,code+'\n'+name,size+3)
 for entries,left in [(ins,True),(outs,False)]:
  for j,(key,label) in enumerate(entries):
   cy=y+(j+.5)*step
   bx=x if left else rx
   nh=min(step-14,size*3.9)
   node(draw,(bx,cy-nh/2,nw,nh),key,size=size-2)
   arrow(draw,(x+nw,cy) if left else (px+pw,cy),(px,cy) if left else (rx,cy),label,size-2)
 return h

# BFD với các nhánh nghiệp vụ, đường nối không mang ý nghĩa dòng dữ liệu.
im=Image.new('RGB',(3600,1220),'white');d=ImageDraw.Draw(im)
node(d,(1120,30,1360,125),'E','HỆ THỐNG MORNING BREW CRM',44)
d.line([(1800,155),(1800,225)],fill='black',width=3)
d.line([(290,225),(3310,225)],fill='black',width=3)
for k in range(1,7):
 x=30+(k-1)*600
 d.line([(x+260,225),(x+260,285)],fill='black',width=3)
 node(d,(x,285,520,130),'E',f'{k}.0\n{P[k]}',34)
 d.line([(x+30,415),(x+30,1055)],fill='black',width=3)
 for j,(name,_,_) in enumerate(C[k],1):
  yy=470+(j-1)*240
  d.line([(x+30,yy+80),(x+85,yy+80)],fill='black',width=3)
  node(d,(x+85,yy,445,160),'E',f'{k}.{j}\n{name}',32)
im.save(QA/'bfd.png')

# Ngữ cảnh: 3 tác nhân, 1 tiến trình, không có kho nội bộ.
im=Image.new('RGB',(2000,1030),'white');d=ImageDraw.Draw(im)
context_i=[('E1','Đặt món, thanh toán, hủy, đổi điểm, phản hồi, trả lời khảo sát'),('E2','Đơn tại quầy, xác nhận thu tiền, tiến độ pha chế và giao hàng'),('E3','Yêu cầu hủy, hoàn, cấp ưu đãi, khảo sát, trả lời và báo cáo')]
context_o=[('E1','Menu, kết quả đơn và tiền, điểm, voucher, phản hồi, khảo sát'),('E2','Menu, đơn cần phục vụ, kết quả lập đơn, thu tiền và giao nhận'),('E3','Kết quả xử lý, phản hồi cần trả lời và báo cáo')]
band(d,30,80,1940,'0','Hệ thống\nMorning Brew CRM',context_i,context_o,step=290,size=36)
im.save(QA/'dfd0.png')

# Toàn bộ mức 1 trên một sơ đồ. Mỗi nhóm là một tiến trình, kho/tác nhân lặp có cùng mã.
im=Image.new('RGB',(3700,2380),'white');d=ImageDraw.Draw(im)
for k in range(1,7):
 x=35+((k-1)%2)*1850;y=40+((k-1)//2)*790
 band(d,x,y,1770,f'{k}.0',P[k],list(SHORT[k]['i'].items()),list(SHORT[k]['o'].items()),step=117,size=28)
im.save(QA/'dfd1.png')

for k in range(1,7):
 heights=[max(len(s[1]),len(s[2]),2)*112 for s in C[k]]
 total=sum(heights)+240
 im=Image.new('RGB',(2200,total),'white');d=ImageDraw.Draw(im);y=30
 for j,(name,ins,outs) in enumerate(C[k],1):
  h=band(d,25,y,2150,f'{k}.{j}',name,ins,outs,step=112,size=31)
  if k==1 and j<3:
   arrow(d,(1100,y+h),(1100,y+h+70))
   text(d,(1130,y+h+5,600,60),'Đơn hợp lệ' if j==1 else 'Đơn đã tính tiền và voucher hợp lệ',26)
  y+=h+70
 im.save(QA/f'dfd2_{k}.png')

doc=Document();s=doc.sections[0]
s.page_width,s.page_height=Cm(42),Cm(29.7)
s.top_margin=s.bottom_margin=Cm(1.3);s.left_margin=s.right_margin=Cm(1.5)
for name in ['Normal','Title','Heading 1','Heading 2']:
 st=doc.styles[name];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0);st.font.size=Pt(12 if name=='Normal' else 20 if name=='Title' else 17)
 for el in list(st.element.iter(qn('w:pBdr'))):el.getparent().remove(el)
doc.styles['Normal'].paragraph_format.space_after=Pt(6)
def para(t):return doc.add_paragraph(t)
def heading(t):doc.add_paragraph(t,'Heading 1')
def picture(path,maxh=22):
 im=Image.open(path);w,h=im.size;scale=min(39/w,maxh/h)
 p=doc.add_paragraph();p.alignment=1;p.paragraph_format.space_after=Pt(2)
 p.add_run().add_picture(str(path),width=Cm(w*scale),height=Cm(h*scale))
def page(t):doc.add_page_break();heading(t)
def table(headers,rows,widths):
 tb=doc.add_table(rows=1,cols=len(headers));tb.autofit=False
 for c,w in zip(tb.columns,widths):c.width=Cm(w)
 for c,t in zip(tb.rows[0].cells,headers):c.text=t
 for row in rows:
  for c,t in zip(tb.add_row().cells,row):c.text=t
 for i,row in enumerate(tb.rows):
  pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
  if i==0:pr.append(OxmlElement('w:tblHeader'))
  for c,w in zip(row.cells,widths):
   c.width=Cm(w);cp=c._tc.get_or_add_tcPr();bo=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');bo.append(el)
   cp.append(bo)
   if i==0:
    el=OxmlElement('w:shd');el.set(qn('w:fill'),'EEEEEE');cp.append(el)
    for r in c.paragraphs[0].runs:r.bold=True
 return tb
doc.add_paragraph('SƠ ĐỒ BFD VÀ DFD HỆ THỐNG MORNING BREW CRM','Title')
para('Nhóm 1 • Báo cáo tuần 4 • Ngày 01 tháng 10 năm 2026')
para('Phạm vi: Các nghiệp vụ đặc trưng phục vụ web Admin, web POS và app khách hàng: đặt món, phục vụ đơn, thanh toán, điểm và voucher, chăm sóc khách hàng, phân tích kinh doanh.')
heading('Quy ước đọc sơ đồ')
para('BFD biểu diễn phân rã chức năng theo cây. DFD mức 0 là sơ đồ ngữ cảnh; mức 1 gồm 6 tiến trình chính; mức 2 phân rã từng tiến trình chính thành 3 tiến trình con. Mã chức năng của BFD được giữ thống nhất với DFD.')
para('Trong DFD: hình chữ nhật là tác nhân ngoài; hình chữ nhật bo góc là tiến trình; hình chữ nhật hở cạnh phải là kho dữ liệu; mũi tên một chiều biểu diễn dữ liệu. Tên trên mũi tên là dữ liệu được truyền, không phải thứ tự thực hiện.')
para('Các tác nhân và kho có cùng mã được vẽ lặp để giảm giao cắt. Ví dụ mọi ký hiệu D3 đều là một kho Đơn và giao hàng. Các tiến trình trao đổi thông tin qua những kho dùng chung này; vị trí các nhóm trên trang không thể hiện trình tự xử lý.')
table(['Mã','Tác nhân','Vai trò'],[
 ('E1','Khách hàng','Đặt và theo dõi đơn trên app; thanh toán, đổi điểm, phản hồi và trả lời khảo sát.'),
 ('E2','Nhân viên POS','Lập đơn tại quầy; nhận đơn app; ghi nhận thu tiền, pha chế, bàn giao và giao hàng.'),
 ('E3','Cấp quản lý','Chủ doanh nghiệp, quản lý khu vực và quản lý chi nhánh; thao tác và xem dữ liệu trong phạm vi được giao.')],[2,7,30])
para('App, POS và Admin là giao diện của cùng hệ thống, không phải ba tác nhân ngoài riêng. Khách vãng lai giao dịch qua nhân viên POS, không bắt buộc có tài khoản. Bộ mô phỏng thanh toán nằm trong hệ thống; chưa có cổng thanh toán thật hoặc tài khoản tài xế.')
heading('Các kho dữ liệu logic')
table(['Mã','Kho dữ liệu','Nội dung chính và bảng tương ứng'],[
 ('D1',N['D1'],'Users, Customers, CustomerAddresses, CustomerPreferences: nhận diện, địa chỉ, sở thích và số dư điểm.'),
 ('D2',N['D2'],'Branches, Products, ProductVariants, các bảng giá và tùy chọn theo chi nhánh: menu, giá, giờ phục vụ, phí giao và tình trạng bán.'),
 ('D3',N['D3'],'Orders, OrderDetails, OrderItemModifiers, Deliveries, OrderStatusHistory: đơn, giá chốt, người nhận và tiến độ xử lý.'),
 ('D4',N['D4'],'LoyaltyPolicies, LoyaltyTransactions, Promotions, CustomerPromotions: chính sách điểm, sổ điểm và vòng đời voucher.'),
 ('D5',N['D5'],'PaymentAttempts, Refunds: từng lần thanh toán, kết quả thu và hoàn tiền.'),
 ('D6',N['D6'],'Feedbacks và các bảng Surveys: phản hồi, trả lời, khảo sát, danh sách nhận và câu trả lời.')],[2,7,30])
para('Kho logic có thể gồm nhiều bảng; đây không phải ERD. Đăng nhập, phân quyền, cấu hình danh mục và lưu nhật ký là chức năng hỗ trợ; không phân rã thành thêm, xóa, sửa trong BFD này. Mọi xử lý vẫn phải kiểm tra quyền và phạm vi dữ liệu.')
page('1 Sơ đồ phân rã chức năng BFD')
picture(QA/'bfd.png',15)
para('Sơ đồ gồm 6 nhóm nghiệp vụ và 18 chức năng con. Đường nối thể hiện quan hệ chức năng cha – con; không dùng mũi tên luồng dữ liệu trong BFD.')
table(['Nhóm','Kết quả nghiệp vụ'],[(f'{k}.0 {P[k]}',t) for k,t in enumerate([
 'Đơn hợp lệ được tính tiền tại server, lưu giá và địa chỉ tại thời điểm đặt.',
 'Đơn được xác nhận, pha chế, giao nhận hoặc hủy theo quyền và trạng thái.',
 'Lưu kết quả thu, hoàn toàn bộ và dữ liệu đối chiếu từng giao dịch.',
 'Cộng hoặc thu hồi điểm, đổi điểm và cấp voucher theo điều kiện chương trình.',
 'Tiếp nhận phản hồi, trả lời chính thức và thu thập khảo sát đúng người nhận.',
 'Báo cáo theo kỳ và phạm vi, từ dữ liệu đơn, tiền và chăm sóc khách hàng.'
],1)],[10,29])
page('2 Sơ đồ luồng dữ liệu mức 0')
picture(QA/'dfd0.png',22)
para('Toàn bộ hệ thống được biểu diễn bằng tiến trình 0. Tác nhân vẽ ở hai phía là bản lặp của cùng E1, E2, E3; mức ngữ cảnh không thể hiện kho dữ liệu nội bộ.')
page('3 Sơ đồ luồng dữ liệu mức 1')
picture(QA/'dfd1.png',23.4)
para('Các gói dữ liệu tổng hợp ở mức 1 được tách thành dữ liệu chi tiết ở mức 2 tương ứng. Các kho D1–D6 và tác nhân E1–E3 dùng chung dù ký hiệu xuất hiện nhiều lần.')
notes={
1:'Một đơn thuộc một chi nhánh. Server kiểm tra menu, size, đường/đá và topping; tính tiền theo giá chi nhánh, mỗi đơn tối đa một voucher. Chốt giá, người nhận và địa chỉ vào đơn. Đơn trả trước chờ thanh toán; đơn trả khi nhận chờ quán xác nhận. Voucher chỉ được giữ cho một đơn tại một thời điểm.',
2:'Đơn trả trước chỉ được đưa vào phục vụ khi thu tiền thành công. Hoàn tất đòi hỏi đã giao/nhận và thu đủ tiền; giao thất bại được ghi riêng để xử lý giao lại hoặc hủy. Khách tự hủy trước xác nhận, sau đó cần quyền quản lý và lý do. Đơn quá hạn thanh toán cũng được xử lý hủy; hủy giải phóng voucher theo hạn gốc. Đơn hủy trong D3 là đầu vào cho hoàn tiền tại 3.2.',
3:'Thanh toán mô phỏng hỗ trợ thành công, thất bại, thử lại và kết quả gửi lặp; mỗi lần thử có bản ghi, mỗi đơn chỉ ghi nhận một lần thu thành công. Chỉ hoàn toàn bộ đơn, tách yêu cầu hoàn khỏi kết quả hoàn. Không coi đơn đã hủy là đã hoàn tiền. Kết quả trong D5 được 2.0 dùng để xử lý đơn và 4.0 dùng để cộng hoặc thu hồi điểm.',
4:'Cộng điểm đúng một lần khi đơn hoàn tất và thu đủ tiền, trên tiền hàng sau giảm, không tính phí giao. Hoàn tiền thu hồi điểm đã cộng; cho phép số dư âm nhưng không đổi voucher nếu thiếu điểm. Đổi điểm và cấp voucher trong cùng giao dịch. Số dư D1 phải khớp sổ điểm D4; chương trình giới hạn thời hạn, phạm vi và số lần cấp.',
5:'Phản hồi gắn đơn phải thuộc khách đó; quản lý chỉ trả lời trong phạm vi được giao, với một trả lời chính thức. Khảo sát chỉ nhận câu trả lời từ khách được mời và còn thời hạn. Mỗi khách trả lời một lần; kiểm tra câu bắt buộc, khoảng điểm và đáp án thuộc câu hỏi. Cấu trúc câu hỏi được khóa sau phát hành.',
6:'Doanh thu ghi nhận khi đơn hoàn tất; chỉ trừ khoản hoàn liên quan doanh thu đã ghi nhận, theo thời điểm hoàn thành hoàn tiền. Phân tích khách và ưu đãi dựa trên lịch sử trong phạm vi quản lý. Báo cáo không bao gồm kế toán, giá vốn hoặc tồn kho nguyên liệu.'
}
for k in range(1,7):
 page(f'4 {k} DFD mức 2 của tiến trình {k} {P[k]}')
 picture(QA/f'dfd2_{k}.png',22.1)
 para('Quy tắc nghiệp vụ: '+notes[k])
page('5 Đối chiếu phân rã và cân bằng dữ liệu')
para('Mỗi tiến trình mức 1 có đúng một sơ đồ mức 2 tương ứng. Những luồng có tên tổng hợp ở mức cha được hiểu là gói chứa các thành phần chi tiết bên dưới; không làm phát sinh tác nhân hoặc kho ngoài phạm vi mức cha.')
table(['Mức 1','Mức 2','Biên dữ liệu giữ nguyên'],[
 (f'{k}.0 {P[k]}',', '.join(f'{k}.{i}' for i in range(1,4)),
 'Đầu vào: '+', '.join(SHORT[k]['i'])+'. Đầu ra: '+', '.join(SHORT[k]['o'])+'.') for k in range(1,7)
],[9,5,25])
heading('Giải thích các gói dữ liệu tổng hợp')
table(['Tiến trình','Thành phần được phân rã ở mức 2'],[
 ('1.0','Yêu cầu đặt món gồm chi nhánh, món, size, tùy chọn và thông tin người nhận. Kết quả gồm menu, lỗi lựa chọn/voucher, chi tiết tiền và mã đơn. Hai luồng nội bộ: 1.1 → 1.2 là đơn hợp lệ; 1.2 → 1.3 là đơn đã tính tiền và voucher hợp lệ.'),
 ('2.0','Dữ liệu xử lý gồm xác nhận quán, tiến độ pha chế, bàn giao, giao lại và yêu cầu hủy. Trạng thái trong D3 bao gồm xác nhận, chuẩn bị, kết quả giao, hoàn tất hoặc hủy; D4 nhận việc sử dụng hoặc giải phóng voucher.'),
 ('3.0','Giao dịch bao gồm lần thử thanh toán, kết quả thu, yêu cầu hoàn và kết quả hoàn. Yêu cầu quản lý gồm hoàn tiền và điều kiện đối chiếu; kết quả quản lý gồm tình trạng hoàn và danh sách giao dịch/sai lệch.'),
 ('4.0','Biến động điểm gồm cộng do mua hàng, trừ khi đổi voucher và thu hồi khi hoàn tiền. Ưu đãi gồm voucher đổi điểm và voucher cấp theo nhóm khách; D4 lưu cả bút toán và danh sách voucher đã cấp.'),
 ('5.0','Dữ liệu chăm sóc gồm phản hồi, trả lời chính thức, nội dung/thời hạn khảo sát, người nhận và bộ câu trả lời. Đầu ra khách gồm xác nhận tiếp nhận, trả lời phản hồi, lời mời khảo sát và kết quả gửi câu trả lời.'),
 ('6.0','Điều kiện báo cáo gồm kỳ, phạm vi và bộ lọc. Báo cáo gồm doanh thu, đơn, món bán chạy, phân tích nhóm khách/ưu đãi và tổng hợp phản hồi/khảo sát.')
],[4,35])
para('Cân bằng ngữ cảnh: mọi luồng với E1, E2, E3 ở mức 1 là thành phần của các gói vào/ra tương ứng trên sơ đồ mức 0. Không bổ sung tác nhân tài xế, ngân hàng hay cổng thanh toán thật trong các mức chi tiết.')
para('Luồng dữ liệu giữa các nhóm qua kho chung: 1.0 ghi đơn vào D3 để 2.0 và 3.0 xử lý; 3.0 ghi kết quả tiền vào D5 để 2.0, 4.0 và 6.0 đọc; 2.0 ghi đơn hoàn tất vào D3 để 4.0 cộng điểm và 6.0 thống kê. Những đường truyền này là dữ liệu, không biểu diễn bằng luồng điều khiển.')
footer=s.footer.paragraphs[0];footer.alignment=2;footer.add_run('Nhóm 1 • Morning Brew CRM • Trang ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
out=BASE/'1.docx';doc.save(out)
(QA/'model.json').write_text(json.dumps({'actors_stores':N,'processes':P,'children':C,'level1':SHORT},ensure_ascii=False,indent=2),encoding='utf-8')
print(out.resolve())
print('6 parent processes / 18 child processes / boundary endpoint checks passed')
