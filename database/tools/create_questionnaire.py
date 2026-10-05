from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

L = 'Likert 5 mức: 1 – Hoàn toàn không cần thiết; 2 – Không cần thiết; 3 – Trung lập; 4 – Cần thiết; 5 – Rất cần thiết.'
rows = [
('Trong một tháng, anh/chị thường mua đồ uống tại chuỗi bao nhiêu lần?', 'Hồ sơ và phân nhóm khách hàng', 'Thứ bậc, chọn một: Chưa mua; 1–3 lần; 4–7 lần; 8–12 lần; Trên 12 lần.'),
('Anh/chị thường sử dụng hình thức nào để mua đồ uống?', 'Kênh bán hàng', 'Định danh, chọn nhiều: Mua trực tiếp tại quầy; Đặt qua điện thoại/tin nhắn; Đặt qua app của quán; Đặt qua ứng dụng giao đồ ăn; Khác.'),
('Anh/chị muốn sử dụng hình thức nhận hàng nào khi đặt qua app?', 'Đặt hàng và giao hàng', 'Định danh, chọn một: Chỉ nhận tại quán; Chỉ giao đến địa chỉ; Cả hai; Chưa có nhu cầu.'),
('Anh/chị đánh giá mức độ cần thiết của việc xem những món đang có bán tại chi nhánh đã chọn?', 'Chi nhánh và thực đơn', L),
('Anh/chị muốn tùy chỉnh những yếu tố nào khi đặt một món?', 'Tùy chọn món', 'Định danh, chọn nhiều: Size; Topping; Lượng đường; Lượng đá; Ghi chú riêng; Không cần tùy chỉnh.'),
('Những thông tin nào về món giúp anh/chị quyết định đặt mua?', 'Thực đơn và chi tiết sản phẩm', 'Định danh, chọn nhiều: Tên món; Hình ảnh; Mô tả; Giá theo size; Phụ phí topping; Tình trạng còn bán.'),
('Anh/chị đánh giá mức độ cần thiết của việc lưu nhiều địa chỉ giao hàng và chọn một địa chỉ mặc định?', 'Địa chỉ khách hàng', L),
('Anh/chị muốn thanh toán đơn hàng vào thời điểm nào?', 'Thanh toán', 'Định danh, chọn một: Thanh toán trước; Thanh toán khi nhận; Muốn có cả hai lựa chọn.'),
('Khi một lần thanh toán thất bại, anh/chị muốn thực hiện thao tác nào tiếp theo?', 'Thanh toán và thử lại', 'Định danh, chọn một: Thử lại trên cùng đơn; Đổi phương thức thanh toán; Hủy đơn; Liên hệ quán hỗ trợ.'),
('Anh/chị đánh giá mức độ cần thiết của việc theo dõi trạng thái đơn trên app?', 'Theo dõi đơn hàng', L),
('Anh/chị muốn nhận thông báo trong app về những nội dung nào?', 'Thông báo', 'Định danh, chọn nhiều: Quán xác nhận đơn; Đơn sẵn sàng; Đơn đang giao; Kết quả thanh toán/hoàn tiền; Voucher được cấp; Lời mời khảo sát; Không muốn nhận.'),
('Anh/chị đánh giá mức độ cần thiết của việc tự hủy đơn trước khi quán xác nhận?', 'Hủy đơn', L),
('Khi đơn đã thanh toán bị hủy, anh/chị muốn theo dõi những thông tin hoàn tiền nào?', 'Hoàn tiền', 'Định danh, chọn nhiều: Số tiền hoàn; Lý do hoàn; Trạng thái xử lý; Thời điểm hoàn thành; Mã tham chiếu giao dịch.'),
('Anh/chị đánh giá mức độ cần thiết của việc xem số dư điểm và lịch sử cộng/trừ điểm?', 'Điểm thành viên', L),
('Anh/chị thích đổi điểm lấy loại voucher nào?', 'Đổi điểm và voucher', 'Định danh, chọn một: Giảm một số tiền cố định; Giảm theo phần trăm; Không có ưu tiên; Không có nhu cầu đổi điểm.'),
('Những điều kiện voucher nào anh/chị cần thấy rõ trước khi sử dụng?', 'Ví voucher', 'Định danh, chọn nhiều: Hạn sử dụng; Đơn tối thiểu; Mức giảm; Trần giảm; Chi nhánh áp dụng; Số điểm cần đổi.'),
('Anh/chị muốn xem những thông tin nào trong lịch sử mua hàng?', 'Lịch sử đơn hàng', 'Định danh, chọn nhiều: Chi nhánh; Món và tùy chọn đã mua; Thời gian đặt; Tổng tiền; Voucher đã dùng; Trạng thái đơn; Trạng thái thanh toán/hoàn tiền.'),
('Anh/chị đánh giá mức độ cần thiết của việc gửi phản hồi gắn với một đơn hàng và xem câu trả lời của quán?', 'Phản hồi và chăm sóc khách hàng', L),
('Anh/chị sẵn sàng dành tối đa bao nhiêu phút để trả lời một khảo sát trên app?', 'Khảo sát khách hàng', 'Tỷ lệ: Nhập số phút từ 0 trở lên: ........ phút. Giá trị 0 nghĩa là không muốn tham gia.'),
('Những bước nào thường gây khó khăn khi lập đơn tại quầy?', 'Lập đơn POS', 'Định danh, chọn nhiều: Tìm món; Chọn size/topping; Nhập đường/đá; Ghi chú; Xác định giá; Áp voucher; Không gặp khó khăn.'),
('Với một đơn thông thường, từ lúc bắt đầu nhập món đến khi chốt đơn, anh/chị hiện mất khoảng bao nhiêu giây?', 'Hiệu quả thao tác POS', 'Tỷ lệ: Nhập số giây ước tính lớn hơn 0: ........ giây. Chọn “Chưa có dữ liệu” nếu chưa thực hiện.'),
('Thông tin nào thuận tiện nhất để anh/chị tìm đúng khách hàng thành viên tại quầy?', 'Nhận diện khách hàng', 'Định danh, chọn một: Số điện thoại; Email; Mã khách hàng; Khác.'),
('Anh/chị cần nhìn thấy những thông tin nào để xử lý đơn từ app?', 'Tiếp nhận và xử lý đơn', 'Định danh, chọn nhiều: Mã đơn; Thời gian đặt; Món và tùy chọn; Ghi chú; Hình thức nhận; Trạng thái thanh toán; Thông tin người nhận.'),
('Khi giao hàng không thành công, anh/chị cần ghi nhận những thông tin nào?', 'Quản lý giao hàng', 'Định danh, chọn nhiều: Lý do thất bại; Thời điểm giao; Người giao; Số điện thoại người giao; Số lần giao; Ghi chú xử lý tiếp theo.'),
('Theo trách nhiệm công việc hiện tại, anh/chị cần xem dữ liệu trong phạm vi nào?', 'Phân quyền và phạm vi dữ liệu', 'Định danh, chọn một: Toàn chuỗi; Các khu vực được phân công; Chi nhánh được phân công.'),
('Những chức năng nào nhân viên POS cần được cấp quyền thực hiện?', 'Phân quyền thao tác', 'Định danh, chọn nhiều: Lập đơn; Xác nhận thanh toán; Cập nhật pha chế; Cập nhật giao hàng; Tra cứu thành viên; Đánh dấu món tạm hết; Khác.'),
('Giá bán của cùng một món cần được quản lý theo cách nào?', 'Giá và thực đơn chi nhánh', 'Định danh, chọn một: Một giá chung toàn chuỗi; Giá riêng từng chi nhánh; Giá chung và cho phép chi nhánh điều chỉnh khi được cấp quyền.'),
('Những tiêu chí nào cần dùng để chọn khách hàng nhận ưu đãi hoặc khảo sát?', 'Phân nhóm khách hàng, ưu đãi và khảo sát', 'Định danh, chọn nhiều: Độ tuổi; Giới tính; Sở thích; Lịch sử mua; Điểm hiện có; Chi nhánh từng mua; Toàn bộ khách hàng.'),
('Anh/chị đánh giá mức độ cần thiết của việc lưu người thực hiện, thời gian và lý do khi hủy đơn hoặc yêu cầu hoàn tiền?', 'Hủy đơn, hoàn tiền và nhật ký thao tác', L),
('Những chỉ tiêu nào anh/chị cần theo dõi trên báo cáo quản trị?', 'Báo cáo và thống kê', 'Định danh, chọn nhiều: Doanh thu sau hoàn tiền; Số đơn hoàn tất; Số đơn hủy; Món bán chạy; Khách quay lại; Voucher đã sử dụng; Điểm đã cấp/thu hồi; Phản hồi chưa xử lý.'),
]
doc = Document()
s = doc.sections[0]
s.page_width, s.page_height = Cm(21), Cm(29.7)
s.top_margin = s.bottom_margin = Cm(1.8)
s.left_margin = s.right_margin = Cm(2)
for name in ['Normal', 'Title', 'Heading 1']:
    style = doc.styles[name]
    style.font.name = 'Times New Roman'
    style.font.color.rgb = RGBColor(0,0,0)
    style.font.size = Pt(12 if name=='Normal' else 15 if name=='Title' else 13)
    for el in list(style.element.iter(qn('w:pBdr'))): el.getparent().remove(el)
normal = doc.styles['Normal'].paragraph_format
normal.line_spacing = 1.05
normal.space_after = Pt(3)
doc.add_paragraph('BẢNG CÂU HỎI KHẢO SÁT PHÂN TÍCH HỆ THỐNG', 'Title')
doc.add_paragraph('Đề tài: Hệ thống quản lý chuỗi quán cà phê Morning Brew CRM').runs[0].bold = True
doc.add_paragraph('Mục đích: Thu thập nhu cầu phục vụ phân tích web Admin, web POS và app khách hàng. Bảng gồm 30 câu hỏi; mỗi câu nêu rõ đối tượng, module, thang đo và phương án trả lời.')
doc.add_paragraph('Hướng dẫn: Trả lời phần phù hợp với vai trò; chọn một hoặc nhiều đáp án theo từng câu. Với lựa chọn “Không có nhu cầu”, “Không muốn nhận” hoặc “Không gặp khó khăn”, không chọn kèm phương án khác. Nếu câu hỏi không thuộc trải nghiệm hoặc trách nhiệm, ghi “Không áp dụng”.')
doc.add_paragraph('Quy ước: Định danh dùng để phân loại; thứ bậc có thứ tự; Likert dùng 5 mức đánh giá; tỷ lệ dùng số lượng hoặc thời gian với giá trị 0 có ý nghĩa.')
for i,(question,module,answer) in enumerate(rows,1):
    if i in [1,20,25]:
        doc.add_paragraph({1:'A Khảo sát khách hàng về app',20:'B Khảo sát nhân viên về web POS',25:'C Khảo sát quản lý về web Admin'}[i],'Heading 1')
    who = 'Khách hàng' if i<20 else 'Nhân viên POS' if i<25 else 'Chủ doanh nghiệp, quản lý chi nhánh' if i==26 else 'Chủ doanh nghiệp, quản lý khu vực, quản lý chi nhánh'
    p=doc.add_paragraph()
    p.paragraph_format.space_before=Pt(7)
    p.add_run(f'Câu {i}. {question}').bold=True
    p.paragraph_format.keep_with_next=True
    for label,text in [('Đối tượng được hỏi: ',who),('Module phục vụ: ',module),('Thang đo và đáp án: ',answer)]:
        p=doc.add_paragraph()
        p.add_run(label).bold=True
        p.add_run(text)
        p.paragraph_format.keep_with_next=label!='Thang đo và đáp án: '
        p.paragraph_format.keep_together=True
footer=s.footer.paragraphs[0]
footer.alignment=2
footer.add_run('Trang ')
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE');footer._p.append(field)
out=Path('database/docs/MorningBrewCRM_Bang_cau_hoi_khao_sat.docx')
doc.save(out)
print(out.resolve())
