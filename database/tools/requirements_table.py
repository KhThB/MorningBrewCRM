from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_ORIENT
import ast
source=ast.parse(Path('database/tools/requirements_questionnaire.py').read_text(encoding='utf-8'))
rows=next(ast.literal_eval(n.value) for n in source.body if isinstance(n,ast.Assign) and any(isinstance(t,ast.Name) and t.id=='rows' for t in n.targets))
doc=Document();s=doc.sections[0];s.orientation=WD_ORIENT.LANDSCAPE
s.page_width,s.page_height=Cm(29.7),Cm(21)
s.top_margin=s.bottom_margin=Cm(1.5);s.left_margin=s.right_margin=Cm(1.5)
for n in ['Normal','Title']:
 st=doc.styles[n];st.font.name='Times New Roman';st.font.size=Pt(11 if n=='Normal' else 15);st.font.color.rgb=RGBColor(0,0,0)
 for el in list(st.element.iter(qn('w:pBdr'))):el.getparent().remove(el)
f=doc.styles['Normal'].paragraph_format;f.space_after=Pt(3);f.line_spacing=1.0
doc.add_paragraph('BẢNG CÂU HỎI KHẢO SÁT YÊU CẦU THIẾT KẾ HỆ THỐNG','Title')
doc.add_paragraph('Đề tài: Hệ thống quản lý chuỗi quán cà phê Morning Brew CRM').runs[0].bold=True
doc.add_paragraph('Mục đích: Làm rõ quy trình nghiệp vụ, dữ liệu, quyền hạn và điều kiện xử lý để thiết kế web Admin, web POS và app khách hàng.')
doc.add_paragraph('Hướng dẫn: Chọn phương án theo quy định của chuỗi và bổ sung chi tiết được hỏi. Đây là phương án để xác nhận yêu cầu, chưa phải kết quả khảo sát. Nếu chưa có quy định, ghi “Cần thống nhất”.')
doc.add_paragraph('Thang đo: Định danh để phân loại; thứ bậc để xác định trình tự hoặc mốc xử lý; tỷ lệ để ghi số lượng, thời gian và giá trị tiền.')
table=doc.add_table(rows=1, cols=5);table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
widths=[0.9,7.1,3.3,3.2,12.2]
for col,w in zip(table.columns,widths):col.width=Cm(w)
for c,t in zip(table.rows[0].cells,['Số','Câu hỏi khai thác yêu cầu','Đối tượng được hỏi','Module phục vụ','Thang đo và phương án trả lời']):
 c.text=t
 for r in c.paragraphs[0].runs:r.bold=True
 shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E7E6E6');c._tc.get_or_add_tcPr().append(shade)
repeat=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(repeat)
for i,(q,who,module,a) in enumerate(rows,1):
 cells=table.add_row().cells
 for c,t in zip(cells,[str(i),q,who,module,a]):c.text=t
for row in table.rows:
 flag=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(flag)
 for j,(c,w) in enumerate(zip(row.cells,widths)):
  c.width=Cm(w);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
  pr=c._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
  for edge in ['top','left','bottom','right']:
   x=OxmlElement('w:'+edge);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'D9D9D9');borders.append(x)
  pr.append(borders);m=OxmlElement('w:tcMar')
  for edge in ['top','left','bottom','right']:
   x=OxmlElement('w:'+edge);x.set(qn('w:w'),'85');x.set(qn('w:type'),'dxa');m.append(x)
  pr.append(m)
  for p in c.paragraphs:
   p.paragraph_format.space_after=Pt(1)
   if j==0:p.alignment=1
footer=s.footer.paragraphs[0];footer.alignment=2;footer.add_run('Trang ')
x=OxmlElement('w:fldSimple');x.set(qn('w:instr'),'PAGE');footer._p.append(x)
doc.save('database/docs/MorningBrewCRM_Bang_cau_hoi_thiet_ke_he_thong.docx')

